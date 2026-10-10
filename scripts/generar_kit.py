"""Genera el producto: producto/Kit-Autonomo-2027.xlsx

Uso: python3 scripts/generar_kit.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

FILAS = 300  # filas de datos en Ingresos y Gastos
SALIDA = Path(__file__).resolve().parent.parent / "producto" / "Kit-Autonomo-2027.xlsx"

AZUL = "1F3A5F"
AZUL_CLARO = "DCE6F2"
VERDE = "2E7D32"
GRIS = "F2F2F2"
EUR = '#,##0.00 "€";[Red]-#,##0.00 "€"'
PCT = "0%"

titulo = Font(bold=True, size=16, color=AZUL)
cab = Font(bold=True, color="FFFFFF")
cab_fill = PatternFill("solid", fgColor=AZUL)
input_fill = PatternFill("solid", fgColor="FFF8E1")
calc_fill = PatternFill("solid", fgColor=GRIS)
borde = Border(bottom=Side(style="thin", color="BFBFBF"))


def cabecera(ws, fila, columnas, anchos):
    for i, (texto, ancho) in enumerate(zip(columnas, anchos), start=1):
        c = ws.cell(row=fila, column=i, value=texto)
        c.font = cab
        c.fill = cab_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = ancho
    ws.row_dimensions[fila].height = 32
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)


wb = Workbook()

# ---------------------------------------------------------------- Instrucciones
ws = wb.active
ws.title = "Instrucciones"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 100
ws["B2"] = "Kit Autónomo 2027 — Facturas, gastos e impuestos en una sola hoja"
ws["B2"].font = titulo
pasos = [
    "CÓMO USARLO (5 minutos):",
    "1. Ve a la hoja «Config» y rellena tus datos, el año y tu % de retención de IRPF (15% general, 7% los 3 primeros años).",
    "2. Cada vez que emitas una factura, añade una fila en «Ingresos». Solo rellenas las celdas amarillas; el resto se calcula solo.",
    "3. Cada vez que pagues un gasto de la actividad, añade una fila en «Gastos» con su categoría y el % deducible.",
    "4. En «Trimestres» tienes calculado el resultado orientativo del Modelo 303 (IVA) y del Modelo 130 (IRPF) de cada trimestre.",
    "5. En «Resumen» ves ingresos, gastos y beneficio mes a mes, con gráfico.",
    "6. En «Factura» escribe el nº de factura y se genera la factura lista para exportar a PDF (Archivo > Exportar > PDF).",
    "",
    "LEYENDA: celdas amarillas = las rellenas tú · celdas grises = fórmulas (no tocar).",
    "",
    "PLAZOS AEAT (303 y 130): T1 hasta el 20 de abril · T2 hasta el 20 de julio · T3 hasta el 20 de octubre · T4 hasta el 30 de enero.",
    "",
    "AVISO: este kit es una herramienta de organización y cálculo orientativo. No sustituye el asesoramiento de un gestor",
    "o asesor fiscal. Casos especiales (recargo de equivalencia, módulos, operaciones intracomunitarias, prorrata) no están cubiertos.",
    "",
    "Funciona en Microsoft Excel, Google Sheets (Archivo > Importar) y LibreOffice Calc.",
]
for i, t in enumerate(pasos, start=4):
    ws.cell(row=i, column=2, value=t).alignment = Alignment(wrap_text=True)
ws["B4"].font = Font(bold=True, size=12)

# ----------------------------------------------------------------------- Config
cfg = wb.create_sheet("Config")
cfg.sheet_view.showGridLines = False
cfg.column_dimensions["A"].width = 34
cfg.column_dimensions["B"].width = 40
cfg["A1"] = "Configuración"
cfg["A1"].font = titulo
datos_cfg = [
    ("Nombre / Razón social", "Tu Nombre Apellido"),
    ("NIF", "00000000T"),
    ("Dirección", "Calle Ejemplo 1, 28001 Madrid"),
    ("Email", "tu@email.com"),
    ("IBAN (para facturas)", "ES00 0000 0000 0000 0000 0000"),
    ("Año fiscal", 2027),
    ("% IVA por defecto", 0.21),
    ("% retención IRPF en facturas", 0.15),
    ("% pago fraccionado Modelo 130", 0.20),
]
for i, (k, v) in enumerate(datos_cfg, start=3):
    cfg.cell(row=i, column=1, value=k).font = Font(bold=True)
    c = cfg.cell(row=i, column=2, value=v)
    c.fill = input_fill
    if isinstance(v, float):
        c.number_format = PCT
# Referencias absolutas
C_ANYO, C_IVA, C_IRPF, C_130 = "Config!$B$8", "Config!$B$9", "Config!$B$10", "Config!$B$11"

cfg["A14"] = "Categorías de gasto"
cfg["A14"].font = Font(bold=True, size=12, color=AZUL)
categorias = [
    "Software y suscripciones", "Material de oficina", "Equipos informáticos",
    "Suministros (luz, internet, móvil)", "Cuota de autónomo (RETA)", "Gestoría / asesoría",
    "Publicidad y marketing", "Formación", "Desplazamientos y dietas", "Alquiler local / coworking",
    "Comisiones bancarias", "Seguros", "Otros",
]
for i, cat in enumerate(categorias, start=15):
    cfg.cell(row=i, column=1, value=cat).fill = input_fill
RANGO_CAT = f"Config!$A$15:$A${14 + len(categorias)}"

# --------------------------------------------------------------------- Ingresos
ing = wb.create_sheet("Ingresos")
cols = ["Fecha", "Nº factura", "Cliente", "NIF cliente", "Concepto", "Base imponible",
        "% IVA", "Cuota IVA", "% IRPF", "Retención IRPF", "Total a cobrar", "Trimestre", "Mes", "¿Cobrada?"]
cabecera(ing, 1, cols, [12, 12, 26, 14, 34, 15, 8, 13, 8, 14, 15, 10, 7, 11])
dv_si = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
ing.add_data_validation(dv_si)
for r in range(2, FILAS + 2):
    for col in "ABCDEFGIN":
        ing[f"{col}{r}"].fill = input_fill
    ing[f"A{r}"].number_format = "DD/MM/YYYY"
    ing[f"F{r}"].number_format = EUR
    ing[f"G{r}"] = f'=IF(F{r}="","",{C_IVA})'
    ing[f"I{r}"] = f'=IF(F{r}="","",{C_IRPF})'
    ing[f"G{r}"].number_format = ing[f"I{r}"].number_format = PCT
    ing[f"H{r}"] = f'=IF(F{r}="","",ROUND(F{r}*G{r},2))'
    ing[f"J{r}"] = f'=IF(F{r}="","",ROUND(F{r}*I{r},2))'
    ing[f"K{r}"] = f'=IF(F{r}="","",F{r}+H{r}-J{r})'
    ing[f"L{r}"] = f'=IF(A{r}="","",ROUNDUP(MONTH(A{r})/3,0))'
    ing[f"M{r}"] = f'=IF(A{r}="","",MONTH(A{r}))'
    for col in "HJKLM":
        ing[f"{col}{r}"].fill = calc_fill
    for col in "HJK":
        ing[f"{col}{r}"].number_format = EUR
    dv_si.add(f"N{r}")
# Fila de ejemplo
ing["A2"] = "=DATE(Config!$B$8,1,15)"
ing["B2"], ing["C2"], ing["D2"], ing["E2"], ing["F2"], ing["N2"] = (
    "2027-001", "Cliente Ejemplo S.L.", "B00000000", "Diseño web (ejemplo, bórrame)", 1000, "Sí")
ing.conditional_formatting.add(
    f"N2:N{FILAS + 1}", CellIsRule(operator="equal", formula=['"No"'], fill=PatternFill("solid", fgColor="FFCDD2")))

# ----------------------------------------------------------------------- Gastos
gas = wb.create_sheet("Gastos")
cols = ["Fecha", "Proveedor", "NIF proveedor", "Concepto", "Categoría", "Base imponible", "% IVA",
        "Cuota IVA", "Total pagado", "% deducible", "Base deducible", "IVA deducible", "Trimestre", "Mes"]
cabecera(gas, 1, cols, [12, 24, 14, 30, 28, 15, 8, 13, 14, 11, 15, 14, 10, 7])
dv_cat = DataValidation(type="list", formula1=RANGO_CAT, allow_blank=True)
gas.add_data_validation(dv_cat)
for r in range(2, FILAS + 2):
    for col in "ABCDEFGJ":
        gas[f"{col}{r}"].fill = input_fill
    gas[f"A{r}"].number_format = "DD/MM/YYYY"
    gas[f"F{r}"].number_format = EUR
    gas[f"G{r}"].number_format = gas[f"J{r}"].number_format = PCT
    gas[f"H{r}"] = f'=IF(F{r}="","",ROUND(F{r}*N(G{r}),2))'
    gas[f"I{r}"] = f'=IF(F{r}="","",F{r}+H{r})'
    gas[f"K{r}"] = f'=IF(F{r}="","",ROUND(F{r}*IF(J{r}="",1,J{r}),2))'
    gas[f"L{r}"] = f'=IF(F{r}="","",ROUND(H{r}*IF(J{r}="",1,J{r}),2))'
    gas[f"M{r}"] = f'=IF(A{r}="","",ROUNDUP(MONTH(A{r})/3,0))'
    gas[f"N{r}"] = f'=IF(A{r}="","",MONTH(A{r}))'
    for col in "HIKLMN":
        gas[f"{col}{r}"].fill = calc_fill
    for col in "HIKL":
        gas[f"{col}{r}"].number_format = EUR
    dv_cat.add(f"E{r}")
gas["A2"] = "=DATE(Config!$B$8,1,20)"
gas["B2"], gas["D2"], gas["E2"], gas["F2"], gas["G2"], gas["J2"] = (
    "Proveedor Ejemplo", "Suscripción software (ejemplo, bórrame)", "Software y suscripciones", 50, 0.21, 1)
gas["A3"] = "=DATE(Config!$B$8,1,31)"
gas["B3"], gas["D3"], gas["E3"], gas["F3"], gas["G3"], gas["J3"] = (
    "Seguridad Social", "Cuota autónomo (ejemplo, bórrame)", "Cuota de autónomo (RETA)", 80, 0, 1)

# ------------------------------------------------------------------- Trimestres
tri = wb.create_sheet("Trimestres")
tri.sheet_view.showGridLines = False
tri["A1"] = "Impuestos trimestrales (orientativo)"
tri["A1"].font = titulo
tri.column_dimensions["A"].width = 52
for col in "BCDEF":
    tri.column_dimensions[col].width = 16
for i, t in enumerate(["T1", "T2", "T3", "T4", "TOTAL AÑO"], start=2):
    c = tri.cell(row=3, column=i, value=t)
    c.font, c.fill, c.alignment = cab, cab_fill, Alignment(horizontal="center")
I_ = f"Ingresos!$F$2:$F${FILAS + 1}"
I_IVA = f"Ingresos!$H$2:$H${FILAS + 1}"
I_RET = f"Ingresos!$J$2:$J${FILAS + 1}"
I_T = f"Ingresos!$L$2:$L${FILAS + 1}"
G_BASE = f"Gastos!$K$2:$K${FILAS + 1}"
G_IVA = f"Gastos!$L$2:$L${FILAS + 1}"
G_T = f"Gastos!$M$2:$M${FILAS + 1}"

filas_tri = [
    ("MODELO 303 — IVA", None, True),
    ("IVA repercutido (cobrado en tus facturas)", lambda q: f"=SUMIFS({I_IVA},{I_T},{q})", False),
    ("IVA soportado deducible (pagado en gastos)", lambda q: f"=SUMIFS({G_IVA},{G_T},{q})", False),
    ("Resultado 303 (positivo = a pagar)", lambda q, c: f"={c}5-{c}6", False),
    ("", None, False),
    ("MODELO 130 — Pago fraccionado IRPF", None, True),
    ("Ingresos del trimestre (base imponible)", lambda q: f"=SUMIFS({I_},{I_T},{q})", False),
    ("Gastos deducibles del trimestre", lambda q: f"=SUMIFS({G_BASE},{G_T},{q})", False),
    ("Ingresos acumulados en el año", "acum10", False),
    ("Gastos acumulados en el año", "acum11", False),
    ("Rendimiento neto acumulado", lambda q, c: f"={c}12-{c}13", False),
    ("Pago fraccionado acumulado (20%)", lambda q, c: f"=MAX(0,{c}14*{C_130})", False),
    ("Retenciones IRPF acumuladas (tus clientes ya pagaron)", "acumret", False),
    ("Pagos 130 de trimestres anteriores", "prev130", False),
    ("Resultado 130 (a ingresar)", lambda q, c: f"=MAX(0,{c}15-{c}16-{c}17)", False),
    ("", None, False),
    ("TOTAL estimado a pagar este trimestre (303 + 130)", lambda q, c: f"=MAX(0,{c}7)+{c}18", False),
    ("Beneficio neto del trimestre (antes de IRPF anual)", lambda q, c: f"={c}10-{c}11", False),
]
cols_q = "BCDE"
for idx, (etq, fx, es_titulo) in enumerate(filas_tri, start=4):
    a = tri.cell(row=idx, column=1, value=etq)
    if es_titulo:
        a.font = Font(bold=True, color=AZUL, size=12)
        continue
    if fx is None:
        continue
    for q, c in enumerate(cols_q, start=1):
        if fx == "acum10":
            f = f"=SUM($B$10:{c}10)"
        elif fx == "acum11":
            f = f"=SUM($B$11:{c}11)"
        elif fx == "acumret":
            f = f'=SUMIFS({I_RET},{I_T},"<="&{q})'
        elif fx == "prev130":
            f = "=0" if q == 1 else f"=SUM($B$18:{cols_q[q - 2]}18)"
        else:
            try:
                f = fx(q)
            except TypeError:
                f = fx(q, c)
        cell = tri[f"{c}{idx}"]
        cell.value = f
        cell.number_format = EUR
        cell.border = borde
    # Columna total del año (solo filas de flujo, no acumulados)
    if idx in (5, 6, 7, 10, 11, 18, 20, 21):
        tri[f"F{idx}"] = f"=SUM(B{idx}:E{idx})"
        tri[f"F{idx}"].number_format = EUR
        tri[f"F{idx}"].font = Font(bold=True)
for c in "BCDEF":
    tri[f"{c}20"].font = Font(bold=True, color="C62828", size=12)
    tri[f"{c}21"].font = Font(bold=True, color=VERDE)
tri["A20"].font = tri["A21"].font = Font(bold=True)
tri["A23"] = ("Consejo: aparta cada mes el IVA cobrado y un 20% del beneficio en una cuenta separada. "
              "Si el 303 sale negativo puedes compensarlo en trimestres siguientes (o pedir devolución en el T4).")
tri["A23"].alignment = Alignment(wrap_text=True)
tri.merge_cells("A23:F24")
tri.row_dimensions[23].height = 30

# ---------------------------------------------------------------------- Resumen
res = wb.create_sheet("Resumen")
res.sheet_view.showGridLines = False
res["A1"] = "Resumen mensual"
res["A1"].font = titulo
for i, (t, w) in enumerate(zip(["Mes", "Ingresos (base)", "Gastos deducibles", "Beneficio", "Pendiente de cobro"],
                               [14, 18, 18, 18, 20]), start=1):
    c = res.cell(row=3, column=i, value=t)
    c.font, c.fill = cab, cab_fill
    res.column_dimensions[c.column_letter].width = w
meses = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto",
         "Septiembre", "Octubre", "Noviembre", "Diciembre"]
I_M = f"Ingresos!$M$2:$M${FILAS + 1}"
I_TOT = f"Ingresos!$K$2:$K${FILAS + 1}"
I_COB = f"Ingresos!$N$2:$N${FILAS + 1}"
G_M = f"Gastos!$N$2:$N${FILAS + 1}"
for m, nombre in enumerate(meses, start=1):
    r = m + 3
    res[f"A{r}"] = nombre
    res[f"B{r}"] = f"=SUMIFS({I_},{I_M},{m})"
    res[f"C{r}"] = f"=SUMIFS({G_BASE},{G_M},{m})"
    res[f"D{r}"] = f"=B{r}-C{r}"
    res[f"E{r}"] = f'=SUMIFS({I_TOT},{I_M},{m},{I_COB},"No")'
    for col in "BCDE":
        res[f"{col}{r}"].number_format = EUR
res["A16"] = "TOTAL"
res["A16"].font = Font(bold=True)
for col in "BCDE":
    res[f"{col}16"] = f"=SUM({col}4:{col}15)"
    res[f"{col}16"].number_format = EUR
    res[f"{col}16"].font = Font(bold=True)
res["A18"], res["B18"] = "Margen de beneficio", "=IF(B16=0,0,D16/B16)"
res["B18"].number_format = PCT
res["A19"], res["B19"] = "Impuestos trimestrales estimados", "=Trimestres!F20"
res["B19"].number_format = EUR
res["A20"], res["B20"] = "Nº de facturas emitidas", f'=COUNTA(Ingresos!$B$2:$B${FILAS + 1})'
for r in (18, 19, 20):
    res[f"A{r}"].font = Font(bold=True)

chart = BarChart()
chart.title = "Ingresos vs gastos por mes"
chart.y_axis.title = "€"
chart.height, chart.width = 8, 18
data = Reference(res, min_col=2, max_col=3, min_row=3, max_row=15)
chart.add_data(data, titles_from_data=True)
chart.set_categories(Reference(res, min_col=1, min_row=4, max_row=15))
res.add_chart(chart, "G3")

# ---------------------------------------------------------------------- Factura
fac = wb.create_sheet("Factura")
fac.sheet_view.showGridLines = False
for col, w in zip("ABCDE", [3, 30, 30, 18, 18]):
    fac.column_dimensions[col].width = w
fac["B2"] = "FACTURA"
fac["B2"].font = Font(bold=True, size=24, color=AZUL)
fac["D2"] = "Nº factura →"
fac["D2"].alignment = Alignment(horizontal="right")
fac["E2"] = "2027-001"
fac["E2"].fill = input_fill
fac["E2"].font = Font(bold=True, size=12)
fila = f"MATCH($E$2,Ingresos!$B$2:$B${FILAS + 1},0)"


def campo(col_ing):
    return f'=IFERROR(INDEX(Ingresos!${col_ing}$2:${col_ing}${FILAS + 1},{fila}),"")'


fac["B4"] = "=Config!B3"
fac["B4"].font = Font(bold=True)
fac["B5"] = '="NIF: "&Config!B4'
fac["B6"] = "=Config!B5"
fac["B7"] = "=Config!B6"
fac["D4"], fac["E4"] = "Fecha:", campo("A")
fac["E4"].number_format = "DD/MM/YYYY"
fac["B9"] = "FACTURAR A:"
fac["B9"].font = Font(bold=True, color=AZUL)
fac["B10"] = campo("C")
fac["B11"] = f'=IFERROR("NIF: "&INDEX(Ingresos!$D$2:$D${FILAS + 1},{fila}),"")'
for i, t in enumerate(["Concepto", "", "", "Importe"], start=2):
    c = fac.cell(row=13, column=i, value=t)
    c.font, c.fill = cab, cab_fill
fac.merge_cells("B14:D14")
fac["B14"] = campo("E")
fac["E14"] = campo("F")
lineas = [
    (16, "Base imponible", campo("F")),
    (17, f'=IFERROR("IVA ("&TEXT(INDEX(Ingresos!$G$2:$G${FILAS + 1},{fila}),"0%")&")","IVA")', campo("H")),
    (18, f'=IFERROR("Retención IRPF (-"&TEXT(INDEX(Ingresos!$I$2:$I${FILAS + 1},{fila}),"0%")&")","IRPF")',
     f'=IFERROR(-INDEX(Ingresos!$J$2:$J${FILAS + 1},{fila}),"")'),
    (19, "TOTAL", campo("K")),
]
for r, etq, f in lineas:
    fac[f"D{r}"] = etq
    fac[f"E{r}"] = f
    fac[f"E{r}"].number_format = EUR
for col in "BCDE":
    fac[f"{col}14"].border = borde
fac["E14"].number_format = EUR
fac["D19"].font = fac["E19"].font = Font(bold=True, size=13)
fac["B22"] = '="Forma de pago: transferencia a "&Config!B7'
fac["B23"] = "Gracias por tu confianza."
fac["B23"].font = Font(italic=True, color="7F7F7F")
fac.print_area = "A1:E24"

SALIDA.parent.mkdir(parents=True, exist_ok=True)
wb.save(SALIDA)
print(f"Generado {SALIDA}")
